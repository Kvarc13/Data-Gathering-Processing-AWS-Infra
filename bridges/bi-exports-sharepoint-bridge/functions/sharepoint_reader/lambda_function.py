# =============================================================================
# File:     lambda_function.py
# Location: bi-automation/lambdas/zeropark/Automations_AWS/bi-exports-sharepoint-bridge/functions/sharepoint_reader/lambda_function.py
# Note:     conceptually "sharepoint-bridge" - the stack keeps its existing
#           deployed name, bi-exports-sharepoint-bridge (see project README).
# =============================================================================
"""
sharepoint-reader - the SharePoint bridge's READ direction.

The bridge already owns everything SharePoint for its WRITE direction
(sharepoint_uploader: the certificate secret, office365, openpyxl, the
cryptography>=42 PEM fix). This function completes the loop: callers invoke
it with a workbook + sheet + S3 destination, and the sheet lands as a CSV.
Data engines never touch SharePoint credentials or its dependencies.

Event contract:
  {
    "file":     "/sites/BITeamCMTech/Shared Documents/.../Some_Input.xlsx",
    "sheet":    "targets",
    "site":     "bicmt",     # optional profile NAME from SITES below
                             # (default: 'bi-team', same site the uploader targets)
    "bucket":   "raw-data-imports-<account>",
    "key":      "MyDataset/2026-08-10/manual_targets.csv"
  }
Returns:
  {"key": "...", "rows": N}     # N = data rows (header excluded)

Site profiles are defined ONCE here, in SITES - callers only name one.
Adding a site = edit this file; every engine pipeline gets it.

Unlike the uploader this function is invoked synchronously and needs no
reserved concurrency or idempotency guard - concurrent reads cannot conflict
and a re-run just re-reads.

connect_to_sharepoint() and _extract_private_key_only() are deliberately the
same code as in sharepoint_uploader: both copies live in THIS stack, so a
certificate rotation still touches exactly one repo folder.
"""

import csv
import io
import json
import logging
import os

import boto3
from office365.sharepoint.client_context import ClientContext
from office365.sharepoint.files.file import File
from openpyxl.reader.excel import load_workbook

logger = logging.getLogger()
logger.setLevel(os.environ.get('LOG_LEVEL', 'INFO'))

s3 = boto3.client('s3', region_name='us-east-1')

# ---------- Site profiles (edit here, in ONE place) ----------
# Callers pick a site by NAME ('site': 'bicmt'); the URL lives only here.

SITES = {
    'bi-team': 'https://centralnic.sharepoint.com/sites/BITeamCMTech',
    'bicmt':   'https://centralnic.sharepoint.com/sites/BICMT',
}
DEFAULT_SITE = 'bi-team'


# ============================================================================
# Lambda entry point  ──  describes the whole job in one screen
# ============================================================================

def lambda_handler(event, context):
    job = parse_job(event)
    logger.info(f"{job['file']} / sheet \"{job['sheet']}\" "
                f"-> s3://{job['bucket']}/{job['key']}")

    credentials = fetch_sharepoint_credentials()
    sharepoint = connect_to_sharepoint(job['site_url'], credentials)
    worksheet = open_workbook_sheet(sharepoint, job['file'], job['sheet'])
    csv_bytes, row_count = worksheet_to_csv(worksheet)
    upload_csv_to_s3(csv_bytes, job['bucket'], job['key'])

    logger.info(f"{row_count} row(s) -> s3://{job['bucket']}/{job['key']}")
    return {'key': job['key'], 'rows': row_count}


# ============================================================================
# Step 0: The job  ──  validate the caller's event, resolve the site
# ============================================================================

def parse_job(event: dict) -> dict:
    missing = [field for field in ('file', 'sheet', 'bucket', 'key')
               if not event.get(field)]
    if missing:
        raise ValueError(f'Missing event field(s): {", ".join(missing)}')

    site = event.get('site') or DEFAULT_SITE
    if site not in SITES:
        raise ValueError(f"Unknown site profile '{site}'. "
                         f"Known: {', '.join(sorted(SITES))}")

    return {**event, 'site_url': SITES[site]}


# ============================================================================
# Step 1: Download the sheet from SharePoint
# ============================================================================

def connect_to_sharepoint(site_url: str, credentials: dict) -> ClientContext:
    cert_settings = dict(credentials['SPCredentials'])  # copy, don't mutate
    cert_settings['private_key'] = _extract_private_key_only(cert_settings['private_key'])
    context = ClientContext(site_url).with_client_certificate('centralnic.com',
                                                              **cert_settings)
    context.load(context.web)
    context.execute_query()
    return context


def _extract_private_key_only(pem_blob: str) -> str:
    """
    The SharePoint secret stores certificate + encrypted private key
    concatenated in a single 'private_key' field. cryptography>=42 doesn't
    tolerate that - it tries to parse the first PEM block (the certificate)
    as the key and fails with InvalidPadding. Extract only the PRIVATE KEY
    block. (Identical to sharepoint_uploader - keep the two in sync.)
    """
    markers = [
        ('-----BEGIN ENCRYPTED PRIVATE KEY-----', '-----END ENCRYPTED PRIVATE KEY-----'),
        ('-----BEGIN PRIVATE KEY-----', '-----END PRIVATE KEY-----'),
        ('-----BEGIN RSA PRIVATE KEY-----', '-----END RSA PRIVATE KEY-----'),
    ]
    for begin, end in markers:
        start_idx = pem_blob.find(begin)
        if start_idx == -1:
            continue
        end_idx = pem_blob.find(end, start_idx)
        if end_idx == -1:
            continue
        return pem_blob[start_idx: end_idx + len(end)]
    raise ValueError('No private key block found in PEM blob')


def open_workbook_sheet(sharepoint: ClientContext, file_path: str, sheet_name: str):
    response = File.open_binary(sharepoint, file_path)
    workbook = load_workbook(io.BytesIO(response.content),
                             read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f'Sheet "{sheet_name}" not found in {file_path} - '
                         f'has: {workbook.sheetnames}')
    return workbook[sheet_name]


# ============================================================================
# Step 2: Convert to CSV and upload to S3
# ============================================================================

def worksheet_to_csv(worksheet) -> tuple[bytes, int]:
    out = io.StringIO()
    writer = csv.writer(out)
    total_rows = 0
    for row in worksheet.iter_rows(values_only=True):
        writer.writerow([clean_cell(value) for value in row])
        total_rows += 1
    return out.getvalue().encode(), max(total_rows - 1, 0)


def clean_cell(value):
    """Excel cells legally contain newlines (Alt+Enter, pasted exports);
    inside a CSV they become multi-line quoted fields that downstream CSV
    sniffers (DuckDB) choke on. Tabular pipelines never need in-cell line
    breaks, so flatten them to spaces at the source - the file this bridge
    emits is then always one line per row."""
    if value is None:
        return ''
    if isinstance(value, str):
        return ' '.join(value.split())   # collapses \r, \n, runs of spaces
    return value


def upload_csv_to_s3(csv_bytes: bytes, bucket: str, key: str) -> None:
    s3.put_object(Bucket=bucket, Key=key, Body=csv_bytes,
                  ContentType='text/csv')


# ============================================================================
# AWS Secrets Manager  (shared utility)
# ============================================================================

def fetch_sharepoint_credentials() -> dict:
    secret_name = os.environ['SHAREPOINT_SECRET_NAME']
    client = boto3.client('secretsmanager', region_name='us-east-1')
    response = client.get_secret_value(SecretId=secret_name)
    return json.loads(response['SecretString'], strict=False)
